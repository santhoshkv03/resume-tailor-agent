"""
Resume Parser utility supporting PDF (.pdf) and Word Document (.docx) parsing,
as well as plain text.
"""
import io
from typing import Union
import pypdf
import docx

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF file bytes."""
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts).strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX file bytes."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also parse tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n".join(paragraphs).strip()

def extract_text(file_bytes: bytes, filename: str) -> str:
    """Detect file type and extract clean text."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        return extract_text_from_docx(file_bytes)
    else:
        # Fallback to UTF-8 decoding
        try:
            return file_bytes.decode("utf-8", errors="ignore").strip()
        except Exception:
            return str(file_bytes)
